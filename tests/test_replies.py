from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from telethon.tl.types import User

from app.ai.contacts import Classified
from app.analytics import build_report, fmt_duration, wilson
from app.jobs.replies import ContactIndex, _scan
from app.store import Store
from app.utils.links import has_link


@pytest.fixture
async def store(tmp_path: Path):
    db = Store(tmp_path / "r.db")
    await db.init()
    return db


# ---------------------------------------------------------------- links


@pytest.mark.parametrize(
    "text",
    [
        "зайди https://x.com",
        "t.me/abc",
        "смотри site.ru",
        "www.a.b",
        "мой сайт example.com/path",
        "tg://resolve?domain=x",
        "bit.ly/abc",
        "telegra.ph/x",
    ],
)
def test_links_detected(text):
    assert has_link(text)


@pytest.mark.parametrize(
    "text",
    ["привет как дела", "Ок.Да", "спасибо, не интересно", "@username", "hello.Thanks", "1.5 млн"],
)
def test_plain_text_is_not_link(text):
    assert not has_link(text)


def test_link_entities_and_preview():
    assert has_link("жми", [{"type": "text_link", "url": "https://a.b"}])
    assert has_link("жми", [SimpleNamespace()]) is False
    assert has_link("", web_preview=True)


# ---------------------------------------------------------------- store + analytics


async def _seed(store: Store):
    acc = await store.add_account("acc1")
    t1 = await store.add_text("offer a", [], title="A")
    t2 = await store.add_text("offer b", [], title="B")
    await store.add_contacts(
        [Classified("username", f"user{i}", 0.9, "t", raw=f"@user{i}", display=f"@user{i}") for i in range(4)]
    )
    contacts = await store.list_contacts(per_page=10)
    sends = {}
    for i, c in enumerate(sorted(contacts, key=lambda c: c.id)):
        text = t1 if i < 2 else t2
        await store.finish_contact(c.id, "sent", account_id=acc.id, text_id=text.id)
        await store.add_send(
            contact_id=c.id,
            account_id=acc.id,
            text_id=text.id,
            proxy_id=None,
            status="sent",
            detail="Основной: ok",
            peer_id=1000 + i,
            message_id=10 + i,
        )
    await store.add_send(
        contact_id=None, account_id=acc.id, text_id=t1.id, proxy_id=None,
        status="error", detail="Основной: Timeout 25s",
    )
    return acc, t1, t2, sorted(contacts, key=lambda c: c.id)


@pytest.mark.asyncio
async def test_reply_dedup_link_filter_and_report(store: Store):
    acc, t1, t2, contacts = await _seed(store)
    idx = await store.reply_contact_index(acc.id)
    assert len(idx) == 4
    ref = next(r for r in idx if r["peer_id"] == 1000)
    now = datetime.now(timezone.utc)

    def add(peer, msg_id, text, link, minutes):
        r = next(x for x in idx if x["peer_id"] == peer)
        return store.add_reply(
            account_id=acc.id, peer_id=peer, tg_msg_id=msg_id, contact_id=r["contact_id"],
            send_id=r["send_id"], text_id=r["text_id"], campaign_id=None, campaign_name="",
            text=text, media="", from_name="N", from_username="u", has_link=link,
            msg_date=(now + timedelta(minutes=minutes)).isoformat(timespec="seconds"),
        )

    first = await add(1000, 1, "привет, интересно", False, -1)
    assert first and first.valid and first.status == "new" and not first.notified
    assert await add(1000, 1, "привет, интересно", False, -1) is None  # дедуп
    redirect = await add(1001, 2, "иди на https://spam.com", True, -2)
    assert redirect and not redirect.valid and redirect.notified  # редирект не уведомляет

    report = await build_report(store)
    ov = report["overview"]
    assert ov["sent"] == 4 and ov["errors"] == 1
    assert ov["responders"] == 1 and ov["reply_rate"] == 25.0
    assert ov["redirect_contacts"] == 1 and ov["redirect_messages"] == 1
    assert ov["reply_messages"] == 1
    text_stats = {x["label"]: x for x in report["dims"]["text"]}
    assert text_stats["A"]["responders"] == 1 and text_stats["B"]["responders"] == 0
    assert report["backlog"]["unanswered"] == 1
    assert report["latency"]["n"] == 1

    answer = await store.add_answer(first, "Здравствуйте!", tg_msg_id=77)
    assert answer.reply_id == first.id
    fresh = await store.get_reply(first.id)
    assert fresh.status == "answered"
    report = await build_report(store)
    assert report["overview"]["answered"] == 1
    assert report["backlog"]["unanswered"] == 0
    assert report["our_response"]["n"] == 1
    thread = await store.reply_thread(fresh)
    assert [t["dir"] for t in thread] == ["in", "out"]
    # фильтр по аккаунту / периоду не падает
    assert (await build_report(store, days=1, account_id=acc.id))["overview"]["sent"] == 4
    assert (await build_report(store, account_id=9999))["overview"]["sent"] == 0


def test_math_helpers():
    lo, hi = wilson(5, 20)
    assert 0 < lo < 25 < hi < 60
    assert wilson(0, 0) == (0.0, 0.0)
    assert fmt_duration(30) == "30 с" and fmt_duration(3700) == "1 ч 1 мин" and fmt_duration(None) == "—"


# ---------------------------------------------------------------- listener scan


class FakeClient:
    def __init__(self, dialogs, messages):
        self._dialogs, self._messages = dialogs, messages

    async def iter_dialogs(self, limit=0, ignore_migrated=True):
        for d in self._dialogs:
            yield d

    async def get_messages(self, entity, limit=0, min_id=0):
        return [m for m in self._messages.get(entity.id, []) if m.id > min_id]


def _msg(mid, text, date, out=False, entities=None):
    return SimpleNamespace(id=mid, raw_text=text, date=date, out=out, entities=entities or [],
                           web_preview=None, action=None)


@pytest.mark.asyncio
async def test_scan_filters_links_and_notifies_once(store: Store):
    from app.jobs import replies as mod

    mod._seen_top.clear()
    acc, *_ = await _seed(store)
    acc = await store.get_account(acc.id)
    now = datetime.now(timezone.utc) + timedelta(minutes=5)

    u0 = User(id=1000, first_name="Ann", username="user0")
    u1 = User(id=1001, first_name="Bob", username="user1")
    stranger = User(id=5555, first_name="X", username="stranger")
    bot = User(id=6666, first_name="Bot", bot=True)
    dialogs = [
        SimpleNamespace(entity=u0, message=SimpleNamespace(id=30)),
        SimpleNamespace(entity=u1, message=SimpleNamespace(id=31)),
        SimpleNamespace(entity=stranger, message=SimpleNamespace(id=32)),
        SimpleNamespace(entity=bot, message=SimpleNamespace(id=33)),
    ]
    messages = {
        1000: [_msg(20, "да, расскажите", now), _msg(21, "я ответил, мой офис site.ru", now), _msg(15, "старое", now, out=True)],
        1001: [_msg(22, "go t.me/scam", now)],
        5555: [_msg(23, "привет", now)],
    }
    client = FakeClient(dialogs, messages)
    idx = ContactIndex(await store.reply_contact_index(acc.id))
    fresh = await _scan(store, acc, client, idx, 100)
    assert [r.tg_msg_id for r in fresh] == [20]  # ссылки не уведомляют, чужие игнорируются
    assert await store.count_replies(valid=True) == 1
    assert await store.count_replies(valid=False) == 2

    again = await _scan(store, acc, client, idx, 100)
    assert again == []  # повторный проход ничего нового


# ---------------------------------------------------------------- screens / xlsx / answer


@pytest.mark.asyncio
async def test_screens_and_xlsx_render(store: Store):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.ui import stats_screens as ui
    from app.utils.report import build_report_xlsx

    acc, t1, t2, contacts = await _seed(store)
    idx = await store.reply_contact_index(acc.id)
    r = idx[0]
    reply = await store.add_reply(
        account_id=acc.id, peer_id=r["peer_id"], tg_msg_id=5, contact_id=r["contact_id"],
        send_id=r["send_id"], text_id=r["text_id"], campaign_id=None, campaign_name="",
        text="<b>да</b> & интересно", media="", from_name="Ann", from_username="ann",
        has_link=False, msg_date=datetime.now(timezone.utc).isoformat(timespec="seconds"),
    )
    report = await build_report(store)
    label = "7 дн"
    for html in (
        ui.overview_html(report, label),
        ui.breakdown_html("Офферы", report["dims"]["text"], label),
        ui.compare_html(report, label, "campaign"),
        ui.time_html(report, label),
        ui.errors_html(report, label),
        ui.reply_notification_html(reply),
        ui.reply_detail_html(reply, await store.reply_thread(reply)),
        ui.replies_list_html([reply], 1, 0, 1),
    ):
        assert 0 < len(html) < 4096
    note = ui.reply_notification_html(reply)
    assert "&lt;b&gt;да&lt;/b&gt; &amp; интересно" in note and "@ann" in note

    data, _ = await build_report_xlsx(store)
    wb = load_workbook(BytesIO(data))
    assert {"Обзор", "Офферы", "Аккаунты", "Ответы", "Редиректы", "По дням"} <= set(wb.sheetnames)
    assert wb["Ответы"].max_row == 2


@pytest.mark.asyncio
async def test_send_answer(store: Store, monkeypatch):
    from contextlib import asynccontextmanager

    from app.jobs import replies as mod

    acc, *_ = await _seed(store)
    await store.update_account(acc.id, telethon_session="/tmp/none.session")
    r = (await store.reply_contact_index(acc.id))[0]
    reply = await store.add_reply(
        account_id=acc.id, peer_id=r["peer_id"], tg_msg_id=9, contact_id=r["contact_id"],
        send_id=r["send_id"], text_id=r["text_id"], campaign_id=None, campaign_name="",
        text="привет", media="", from_name="A", from_username="a", has_link=False,
        msg_date=datetime.now(timezone.utc).isoformat(timespec="seconds"),
    )
    redirect = await store.add_reply(
        account_id=acc.id, peer_id=r["peer_id"], tg_msg_id=10, contact_id=r["contact_id"],
        send_id=r["send_id"], text_id=r["text_id"], campaign_id=None, campaign_name="",
        text="t.me/x", media="", from_name="A", from_username="a", has_link=True,
        msg_date=datetime.now(timezone.utc).isoformat(timespec="seconds"),
    )

    calls = []

    class _Action:
        async def __aenter__(self): return self
        async def __aexit__(self, *a): return False

    class C:
        async def get_input_entity(self, pid): return SimpleNamespace(id=pid)
        def action(self, *a, **k): return _Action()
        async def send_message(self, entity, text, **kw):
            calls.append((entity.id, text))
            return SimpleNamespace(id=500)
        async def send_read_acknowledge(self, entity): pass

    @asynccontextmanager
    async def fake_client(session, proxy=None):
        yield C()

    async def no_sleep(*_a, **_k): pass

    monkeypatch.setattr(mod, "telethon_client", fake_client)
    monkeypatch.setattr(mod.asyncio, "sleep", no_sleep)

    with pytest.raises(mod.AnswerError):
        await mod.send_answer(store, redirect.id, "hi")  # на редирект не отвечаем
    with pytest.raises(mod.AnswerError):
        await mod.send_answer(store, reply.id, "   ")

    ans = await mod.send_answer(store, reply.id, "Здравствуйте", source="api")
    assert calls == [(r["peer_id"], "Здравствуйте")]
    assert ans.tg_msg_id == 500 and ans.source == "api"
    assert (await store.get_reply(reply.id)).status == "answered"


@pytest.mark.asyncio
async def test_flush_notifications_retries(store: Store):
    from app.jobs.replies import flush_notifications

    acc, *_ = await _seed(store)
    r = (await store.reply_contact_index(acc.id))[0]
    reply = await store.add_reply(
        account_id=acc.id, peer_id=r["peer_id"], tg_msg_id=1, contact_id=r["contact_id"],
        send_id=r["send_id"], text_id=r["text_id"], campaign_id=None, campaign_name="",
        text="да", media="", from_name="A", from_username="a", has_link=False,
        msg_date=datetime.now(timezone.utc).isoformat(timespec="seconds"),
    )
    seen = []

    async def failing(rep): raise RuntimeError("no chat")

    async def ok(rep): seen.append(rep.id)

    assert await flush_notifications(store, failing) == 0
    assert (await store.get_reply(reply.id)).notified == 0  # осталось в очереди
    assert await flush_notifications(store, ok) == 1
    assert seen == [reply.id]
    assert await flush_notifications(store, ok) == 0  # второй раз не шлём
