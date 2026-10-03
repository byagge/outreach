from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.analytics import fmt_duration
from app.models import Reply

_DIM_SHEETS = [
    ("text", "Офферы"),
    ("account", "Аккаунты"),
    ("campaign", "Кампании"),
    ("base", "Базы"),
    ("kind", "Тип контакта"),
    ("proxy", "Прокси"),
    ("hour", "Часы"),
    ("weekday", "Дни недели"),
]

_DIM_COLS = [
    ("label", "Название"),
    ("attempts", "Попыток"),
    ("sent", "Отправлено"),
    ("errors", "Ошибок"),
    ("skipped", "Пропусков"),
    ("delivery_rate", "Доставка, %"),
    ("responders", "Ответили"),
    ("reply_rate", "Ответов, %"),
    ("ci_lo", "95% ДИ от"),
    ("ci_hi", "95% ДИ до"),
    ("reply_messages", "Сообщений-ответов"),
    ("redirect_contacts", "Редиректов (ссылки)"),
    ("answered", "Мы ответили"),
    ("unanswered", "Без ответа"),
    ("median_reply", "Медиана до ответа"),
]


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, col in enumerate(header, start=1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    ws.freeze_panes = "A2"


def _reply_rows(replies: list[Reply]) -> list[list[Any]]:
    return [
        [
            r.id,
            (r.msg_date or "")[:19].replace("T", " "),
            r.who,
            r.account_label,
            r.campaign_name,
            r.text_title,
            r.text or (f"[{r.media}]" if r.media else ""),
            {"new": "ждёт ответа", "answered": "отвечено", "ignored": "пропущено"}.get(r.status, r.status),
        ]
        for r in replies
    ]


def report_to_xlsx(
    report: dict[str, Any],
    replies: list[Reply] | None = None,
    redirects: list[Reply] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Обзор"
    ov = report["overview"]
    period = report["period"]
    pairs: list[tuple[str, Any]] = [
        ("Период с", period.get("since") or "начало"),
        ("Период по", period.get("until") or "сейчас"),
        ("Часовой пояс", period.get("timezone")),
        ("Попыток отправки", ov["attempts"]),
        ("Отправлено", ov["sent"]),
        ("Ошибок", ov["errors"]),
        ("Пропусков", ov["skipped"]),
        ("Доставка, %", ov["delivery_rate"]),
        ("Уникальных контактов", ov["contacts_reached"]),
        ("Ответили (люди)", ov["responders"]),
        ("Ответов, % от отправленных", ov["reply_rate"]),
        ("95% ДИ", f"{ov['reply_rate_ci'][0]}–{ov['reply_rate_ci'][1]}"),
        ("Сообщений-ответов", ov["reply_messages"]),
        ("Редиректов (ответ со ссылкой, не засчитан)", ov["redirect_contacts"]),
        ("Мы ответили", ov["answered"]),
        ("Диалог продолжился", ov["continued"]),
        ("Ждут ответа сейчас", report["backlog"]["unanswered"]),
        ("Время до ответа: медиана", fmt_duration(report["latency"]["median"])),
        ("Время до ответа: среднее", fmt_duration(report["latency"]["avg"])),
        ("Время до ответа: p90", fmt_duration(report["latency"]["p90"])),
        ("Наша реакция: медиана", fmt_duration(report["our_response"]["median"])),
    ]
    ws.append(["Показатель", "Значение"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for k, v in pairs:
        ws.append([k, v])
    ws.append([])
    ws.append(["Выводы"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for text in report["insights"]:
        ws.append([text])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 28

    _sheet(
        wb,
        "Воронка",
        ["Этап", "Кол-во", "% от предыдущего", "% от отправленных"],
        [[f["stage"], f["count"], f["pct_of_prev"], f["pct_of_sent"]] for f in report["funnel"]],
    )
    _sheet(
        wb,
        "Скорость ответа",
        ["Интервал", "Ответов", "%"],
        [[b["label"], b["count"], b["pct"]] for b in report["latency"]["buckets"]],
    )

    for key, title in _DIM_SHEETS:
        rows = []
        for x in report["dims"][key]:
            data = dict(x)
            data["ci_lo"], data["ci_hi"] = x["reply_rate_ci"]
            data["median_reply"] = fmt_duration(x["median_reply_sec"])
            rows.append([data[c] for c, _ in _DIM_COLS])
        _sheet(wb, title, [h for _, h in _DIM_COLS], rows)

    _sheet(
        wb,
        "По дням",
        ["Дата", "Отправлено", "Ошибок", "Ответили (по отправкам дня)", "Ответов, %",
         "Ответов получено", "Редиректов получено"],
        [
            [d["date"], d["sent"], d["errors"], d["responders"], d["reply_rate"],
             d["replies_received"], d["redirects_received"]]
            for d in report["daily"]
        ],
    )
    _sheet(
        wb,
        "Когда отвечают",
        ["Час", "Ответов"],
        [[h["label"], h["replies"]] for h in report["reply_hours"]],
    )
    _sheet(
        wb,
        "Ошибки",
        ["Статус", "Причина", "Кол-во"],
        [[e["status"], e["reason"], e["count"]] for e in report["errors"]],
    )
    reply_header = ["ID", "Когда (UTC)", "Кто", "Аккаунт", "Кампания", "Оффер", "Текст", "Статус"]
    _sheet(wb, "Ответы", reply_header, _reply_rows(replies or []))
    _sheet(wb, "Редиректы", reply_header, _reply_rows(redirects or []))

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


async def build_report_xlsx(
    store,
    *,
    days: int | None = None,
    since: str | None = None,
    until: str | None = None,
    campaign_id: int | None = None,
    account_id: int | None = None,
) -> tuple[bytes, dict[str, Any]]:
    """Отчёт целиком: считает аналитику, подтягивает ответы/редиректы, пишет XLSX."""
    from app.analytics import build_report, period_bounds

    report = await build_report(
        store, days=days, since=since, until=until, campaign_id=campaign_id, account_id=account_id
    )
    start, end = period_bounds(days, since, until)
    kw = dict(since=start, until=end, campaign_id=campaign_id, account_id=account_id, limit=5000)
    replies = await store.list_replies(valid=True, **kw)
    redirects = await store.list_replies(valid=False, **kw)
    return report_to_xlsx(report, replies, redirects), report
