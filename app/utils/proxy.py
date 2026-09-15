from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

from openpyxl import load_workbook

SCHEME_ALIASES: dict[str, str] = {
    "socks5": "socks5",
    "socks": "socks5",
    "sock": "socks5",
    "s5": "socks5",
    "socks5h": "socks5",
    "socks5-hostname": "socks5",
    "socks5hostname": "socks5",
    "socks4": "socks4",
    "socks4a": "socks4",
    "s4": "socks4",
    "http": "http",
    "https": "https",
    "ssl": "https",
    "tls": "https",
    "http-connect": "http",
    "connect": "http",
    "mtproto": "mtproto",
    "mtproxy": "mtproto",
    "mtp": "mtproto",
    "tg": "mtproto",
}

_DEFAULT_PORTS = {
    "socks5": 1080,
    "socks4": 1080,
    "http": 80,
    "https": 443,
    "mtproto": 443,
}

_HTTP_PORTS = {
    80,
    81,
    8000,
    8001,
    8008,
    8080,
    8081,
    8088,
    8118,
    8123,
    8888,
    8889,
    3128,
    3129,
    3127,
    9090,
    9091,
    10000,
}
_HTTPS_PORTS = {443, 444, 8443, 4443, 9443, 10443}
_SOCKS_PORTS = {1080, 1081, 1082, 1085, 1088, 4145, 4153, 9050, 9051, 9150, 7777, 10808}


@dataclass
class ParsedProxy:
    scheme: str
    host: str
    port: int
    username: str = ""
    password: str = ""
    raw: str = ""
    explicit: bool = False


def detect_scheme(token: str | None) -> str | None:
    raw = re.sub(r"[^a-z0-9\-]+", "", (token or "").strip().lower())
    if not raw:
        return None
    return SCHEME_ALIASES.get(raw)


def guess_scheme_by_port(port: int) -> str:
    if port in _HTTPS_PORTS:
        return "https"
    if port in _HTTP_PORTS:
        return "http"
    if port in _SOCKS_PORTS:
        return "socks5"
    return "socks5"


def _clean(line: str) -> str:
    text = (line or "").strip()
    if not text or text.startswith("#") or text.startswith("//"):
        return ""
    text = re.sub(r"^(?:proxy|proxies)\s*[:=]\s*", "", text, flags=re.I)
    return text.strip().strip("`'\"")


def _looks_host(value: str) -> bool:
    value = (value or "").strip().strip("[]")
    if not value:
        return False
    if re.fullmatch(r"\d{1,3}(?:\.\d{1,3}){3}", value):
        return True
    if ":" in value:
        return True
    if "." in value:
        return True
    return bool(re.fullmatch(r"[A-Za-z0-9._-]+", value)) and not value.isdigit()


def _normalize_separators(text: str) -> str:
    if "://" in text or "@" in text:
        return text
    if "\t" in text:
        return text.replace("\t", ":")
    if ";" in text and text.count(":") <= 1:
        return text.replace(";", ":")
    if text.count(",") >= 2 and "://" not in text:
        return text.replace(",", ":")
    return text


def _parse_mtproto(text: str) -> ParsedProxy | None:
    lowered = text.lower()
    if "t.me/proxy" not in lowered and "t.me/socks" not in lowered and "server=" not in lowered:
        if not lowered.startswith("mtproto") and not lowered.startswith("tg://"):
            return None
    if "://" not in text and "server=" in lowered:
        text = "https://t.me/proxy?" + text.lstrip("?&")
    parsed = urlparse(text.replace("tg://", "https://", 1))
    qs = parse_qs(parsed.query)
    host = (qs.get("server") or qs.get("ip") or [parsed.hostname or ""])[0]
    port_s = (qs.get("port") or [str(parsed.port or "")])[0]
    secret = (qs.get("secret") or [""])[0]
    if not host:
        return None
    port = int(port_s) if str(port_s).isdigit() else 443
    return ParsedProxy("mtproto", host.strip(), port, "", secret, text, explicit=True)


def _from_url(text: str) -> ParsedProxy | None:
    parsed = urlparse(text)
    scheme = detect_scheme(parsed.scheme)
    if parsed.scheme.lower() in {"tg"} or "secret=" in (parsed.query or ""):
        mt = _parse_mtproto(text)
        if mt:
            return mt
    host = parsed.hostname or ""
    user = unquote(parsed.username or "")
    password = unquote(parsed.password or "")
    if parsed.query and not password:
        qs = parse_qs(parsed.query)
        user = user or (qs.get("user") or qs.get("username") or [""])[0]
        password = (qs.get("pass") or qs.get("password") or qs.get("secret") or [""])[0]
    if not host:
        return None
    port = int(parsed.port or 0)
    if not scheme:
        scheme = guess_scheme_by_port(port or 0) if port else "socks5"
        explicit = False
    else:
        explicit = True
    if not port:
        port = _DEFAULT_PORTS.get(scheme, 1080)
    return ParsedProxy(scheme, host, port, user, password, text, explicit=explicit)


def _split_host_port_auth(parts: list[str]) -> tuple[str, str, str, str] | None:
    if len(parts) < 2:
        return None
    if parts[1].isdigit() and _looks_host(parts[0]):
        host, port_s = parts[0], parts[1]
        user = parts[2] if len(parts) >= 3 else ""
        password = ":".join(parts[3:]) if len(parts) >= 4 else ""
        return host, port_s, user, password
    if len(parts) >= 4 and parts[3].isdigit() and _looks_host(parts[2]):
        return parts[2], parts[3], parts[0], parts[1]
    return None


def parse_proxy_line(line: str) -> ParsedProxy | None:
    text = _clean(line)
    if not text:
        return None
    raw = text
    mt = _parse_mtproto(text)
    if mt:
        mt.raw = raw
        return mt

    if "://" in text:
        got = _from_url(text)
        if got:
            got.raw = raw
            return got
        text = text.split("://", 1)[1]

    scheme: str | None = None
    explicit = False
    tokens = [t for t in re.split(r"\s+", text) if t]
    if len(tokens) >= 2 and detect_scheme(tokens[0]):
        scheme = detect_scheme(tokens[0])
        explicit = True
        text = " ".join(tokens[1:])
        if len(tokens) >= 3 and tokens[2].isdigit() and ":" not in tokens[1]:
            user = tokens[3] if len(tokens) >= 4 else ""
            password = " ".join(tokens[4:]) if len(tokens) >= 5 else ""
            host = tokens[1].strip("[]")
            port = int(tokens[2])
            if host and 0 < port <= 65535:
                return ParsedProxy(scheme or "socks5", host, port, user, password, raw, True)

    text = _normalize_separators(text)

    if "@" in text:
        left, right = text.rsplit("@", 1)
        left_parts = left.split(":")
        right_parts = right.split(":")
        if len(left_parts) >= 2 and left_parts[-1].isdigit() and _looks_host(left_parts[0].strip("[]")):
            host = ":".join(left_parts[:-1]).strip("[]")
            port_s = left_parts[-1]
            user, _, password = right.partition(":")
            if detect_scheme(password):
                scheme = detect_scheme(password)
                explicit = True
                password = ""
            elif detect_scheme(right_parts[-1]) and len(right_parts) >= 2:
                scheme = detect_scheme(right_parts[-1])
                explicit = True
                user, password = right_parts[0], ":".join(right_parts[1:-1])
        else:
            user, _, password = left.partition(":")
            if detect_scheme(right_parts[-1]) and not right_parts[-1].isdigit():
                scheme = detect_scheme(right_parts[-1])
                explicit = True
                right_parts = right_parts[:-1]
            host = right_parts[0].strip("[]") if right_parts else ""
            port_s = right_parts[1] if len(right_parts) >= 2 else ""
        if host and port_s.isdigit():
            port = int(port_s)
            if not scheme:
                scheme = guess_scheme_by_port(port)
            return ParsedProxy(scheme, host, port, unquote(user), unquote(password), raw, explicit)

    parts = text.split(":")
    if parts and detect_scheme(parts[0]):
        scheme = detect_scheme(parts[0])
        explicit = True
        parts = parts[1:]
    if parts and detect_scheme(parts[-1]):
        scheme = scheme or detect_scheme(parts[-1])
        explicit = True
        parts = parts[:-1]

    split = _split_host_port_auth(parts)
    if not split:
        return None
    host, port_s, user, password = split
    host = host.strip().strip("[]")
    if not host or not port_s.isdigit():
        return None
    port = int(port_s)
    if port <= 0 or port > 65535:
        return None
    if detect_scheme(password) and ":" not in password:
        scheme = scheme or detect_scheme(password)
        explicit = True
        password = ""
    if not scheme:
        scheme = guess_scheme_by_port(port)
        explicit = False
    return ParsedProxy(scheme, host, port, unquote(user), unquote(password), raw, explicit)


def parse_proxy_blob(blob: str) -> list[ParsedProxy]:
    out: list[ParsedProxy] = []
    seen: set[tuple[str, str, int, str]] = set()
    for line in (blob or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        item = parse_proxy_line(line)
        if not item:
            continue
        key = (item.scheme, item.host, item.port, item.username)
        if key in seen:
            continue
        seen.add(key)
        out.append(item)
    return out


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _from_table(rows: list[list[str]]) -> list[ParsedProxy]:
    if not rows:
        return []
    header = [re.sub(r"[^a-z0-9]+", "", (c or "").strip().lower()) for c in rows[0]]

    def col(*names: str) -> int | None:
        for i, h in enumerate(header):
            if any(n in h for n in names):
                return i
        return None

    type_i = col("type", "scheme", "proto", "protocol", "тип")
    host_i = col("host", "ip", "addr", "address", "server", "хост")
    port_i = col("port", "порт")
    user_i = col("user", "login", "username", "логин")
    pass_i = col("pass", "password", "pwd", "secret", "пароль")
    has_header = host_i is not None and port_i is not None
    data_rows = rows[1:] if has_header else rows
    blob_lines: list[str] = []
    for row in data_rows:
        cells = [(c or "").strip() for c in row]
        if not any(cells):
            continue
        if has_header:
            host = cells[host_i] if host_i is not None and host_i < len(cells) else ""
            port = cells[port_i] if port_i is not None and port_i < len(cells) else ""
            user = cells[user_i] if user_i is not None and user_i < len(cells) else ""
            password = cells[pass_i] if pass_i is not None and pass_i < len(cells) else ""
            scheme = cells[type_i] if type_i is not None and type_i < len(cells) else ""
            bits = [scheme, host, port, user, password]
            blob_lines.append(":".join(b for b in bits if b))
        else:
            blob_lines.append(":".join(c for c in cells if c))
    return parse_proxy_blob("\n".join(blob_lines))


def parse_proxy_file(data: bytes, filename: str = "") -> list[ParsedProxy]:
    name = (filename or "").lower()
    ext = Path(name).suffix.lower()
    if ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = wb.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append(cells)
        wb.close()
        return _from_table(rows)
    text = _decode(data)
    stripped = text.lstrip()
    if ext == ".json" or stripped[:1] in {"{", "["}:
        try:
            payload = json.loads(stripped)
        except json.JSONDecodeError:
            payload = None
        if isinstance(payload, list):
            lines = []
            for item in payload:
                if isinstance(item, str):
                    lines.append(item)
                elif isinstance(item, dict):
                    bits = [
                        str(item.get(k, "") or "")
                        for k in (
                            "type",
                            "scheme",
                            "protocol",
                            "host",
                            "ip",
                            "port",
                            "username",
                            "user",
                            "password",
                            "pass",
                        )
                        if item.get(k) not in (None, "")
                    ]
                    lines.append(":".join(bits) if bits else json.dumps(item))
            return parse_proxy_blob("\n".join(lines))
        if isinstance(payload, dict):
            inner = payload.get("proxies") or payload.get("items") or payload.get("data")
            if isinstance(inner, list):
                return parse_proxy_file(json.dumps(inner).encode("utf-8"), "proxies.json")
    if ext in {".csv", ".tsv"} or (text.count(",") >= 2 and "\n" in text):
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
            if ext == ".tsv":
                dialect.delimiter = "\t"
        rows = [[(c or "").strip() for c in row] for row in csv.reader(StringIO(text), dialect)]
        parsed = _from_table(rows)
        if parsed:
            return parsed
    return parse_proxy_blob(text)
