from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook

USERNAME_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_]{4,31}$")
AT_RE = re.compile(r"^@([A-Za-z][A-Za-z0-9_]{4,31})$")
TME_RE = re.compile(
    r"(?:https?://)?(?:t\.me|telegram\.me|telegram\.dog)/@?([A-Za-z][A-Za-z0-9_]{4,31})(?:/\S*)?$",
    re.I,
)
PHONE_RE = re.compile(r"^\+?[1-9]\d{6,14}$")
ID_RE = re.compile(r"^-?\d{5,18}$")
MD_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BULLET_RE = re.compile(r"^[\s>*\-•–—]+")
NUMBERED_RE = re.compile(r"^\d{1,3}[.)]\s+")

SKIP_HEADERS = {
    "username",
    "user",
    "usernames",
    "telegram",
    "tg",
    "phone",
    "phones",
    "mobile",
    "номер",
    "телефон",
    "id",
    "user_id",
    "userid",
    "tg_id",
    "chat_id",
    "contact",
    "контак",
    "name",
    "имя",
    "first_name",
    "last_name",
    "comment",
    "note",
    "email",
    "mail",
}

CONTACT_HEADERS = {
    "username": "username",
    "user": "username",
    "usernames": "username",
    "telegram": "username",
    "tg": "username",
    "login": "username",
    "ник": "username",
    "никнейм": "username",
    "phone": "phone",
    "phones": "phone",
    "mobile": "phone",
    "tel": "phone",
    "номер": "phone",
    "телефон": "phone",
    "id": "user_id",
    "user_id": "user_id",
    "userid": "user_id",
    "tg_id": "user_id",
    "telegram_id": "user_id",
    "chat_id": "user_id",
    "contact": "any",
    "контак": "any",
}


@dataclass
class Classified:
    kind: str
    value: str
    confidence: float
    reason: str
    raw: str = ""
    extra: str = ""
    display: str = ""


@dataclass
class ParseResult:
    contacts: list[Classified] = field(default_factory=list)
    skipped: int = 0
    notes: list[str] = field(default_factory=list)


def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _digits_phone(token: str) -> str:
    token = token.strip().replace(" ", "").replace("-", "").replace("(", "").replace(")", "")
    if token.startswith("00") and len(token) > 8:
        token = "+" + token[2:]
    if token.startswith("8") and len(token) == 11:
        token = "+7" + token[1:]
    return token


def classify_token(token: str) -> Classified | None:
    raw = (token or "").strip()
    if not raw:
        return None
    text = raw.strip("`'\"[]()<>.,;")
    text = BULLET_RE.sub("", text).strip()
    text = NUMBERED_RE.sub("", text).strip()
    if not text or text.lower() in SKIP_HEADERS:
        return None
    if text.startswith("#") or text.startswith("//"):
        return None

    md = MD_LINK_RE.search(text)
    if md:
        inner = classify_token(md.group(2)) or classify_token(md.group(1))
        if inner:
            inner.raw = raw
            return inner

    tme = TME_RE.search(text.replace(" ", ""))
    if tme:
        uname = tme.group(1)
        if uname.lower() not in {"joinchat", "addstickers", "share", "s", "c", "iv"}:
            return Classified(
                "username",
                uname.lower(),
                0.97,
                "ссылка t.me",
                raw=raw,
                display=f"@{uname}",
            )

    at = AT_RE.match(text) or (USERNAME_RE.match(text) if text[:1].isalpha() else None)
    if at and (text.startswith("@") or (text.isascii() and "_" in text) or text[:1].isalpha()):
        uname = at.group(1) if hasattr(at, "group") and at.lastindex else text.lstrip("@")
        if text.startswith("@") or USERNAME_RE.match(text):
            if text.startswith("@") or (len(text) >= 5 and not text.isdigit()):
                if USERNAME_RE.match(uname) and not uname.isdigit():
                    if text.startswith("@"):
                        conf = 0.95
                    elif "_" in uname or any(ch.isdigit() for ch in uname):
                        conf = 0.8
                    else:
                        conf = 0.55
                    return Classified(
                        "username",
                        uname.lower(),
                        conf,
                        "username" if text.startswith("@") else "похоже на username",
                        raw=raw,
                        display=f"@{uname}",
                    )

    phone_src = _digits_phone(text)
    phone_digits = phone_src.lstrip("+")
    looks_phone = bool(
        text.strip().startswith("+")
        or text.strip().startswith("00")
        or any(ch in raw for ch in " -()")
        or (len(phone_digits) == 11 and phone_digits[0] in "78")
        or (len(phone_digits) >= 12 and PHONE_RE.match(phone_src.replace(" ", "")))
    )
    if looks_phone and PHONE_RE.match(phone_src.replace(" ", "")):
        value = phone_src if phone_src.startswith("+") else f"+{phone_src}"
        conf = 0.93 if text.strip().startswith("+") else 0.82
        return Classified("phone", value, conf, "номер телефона", raw=raw, display=value)

    id_src = text.replace(" ", "")
    if ID_RE.match(id_src) and not (id_src.startswith("0") and len(id_src) > 1):
        if id_src.startswith("-100") and len(id_src) >= 10:
            return Classified(
                "user_id",
                id_src,
                0.6,
                "id канала/чата — для лички сомнительно",
                raw=raw,
                display=id_src,
            )
        conf = 0.9 if len(id_src) >= 8 else 0.75
        return Classified("user_id", id_src, conf, "числовой Telegram ID", raw=raw, display=id_src)

    if USERNAME_RE.match(text) and not text.isdigit():
        return Classified(
            "username",
            text.lower(),
            0.55,
            "слово похоже на username",
            raw=raw,
            display=f"@{text}",
        )

    return Classified("unknown", text, 0.15, "не распознал тип", raw=raw, display=text)


def classify_line(line: str) -> Classified | None:
    raw = (line or "").strip()
    if not raw:
        return None
    text = BULLET_RE.sub("", raw).strip()
    text = NUMBERED_RE.sub("", text).strip()
    if not text or text.startswith("#"):
        return None

    candidates: list[Classified] = []
    parts = re.split(r"[\s,;|/]+", text)
    direct = classify_token(text)
    if direct:
        candidates.append(direct)
    if len(parts) > 1:
        for part in parts:
            item = classify_token(part)
            if item:
                candidates.append(item)
    if not candidates:
        return None
    best = max(candidates, key=lambda c: (c.confidence, 1 if c.kind != "unknown" else 0))
    if best.kind == "unknown" and best.confidence < 0.4:
        return None
    if (
        best.kind == "username"
        and best.confidence < 0.7
        and not raw.strip().startswith("@")
        and "t.me/" not in raw.lower()
        and "_" not in best.value
        and not any(ch.isdigit() for ch in best.value)
    ):
        return None
    best.raw = raw
    leftover = [p for p in parts if p and p.lower().lstrip("@") != best.value.lower().lstrip("@+")]
    if leftover:
        best.extra = " ".join(leftover)[:240]
    return best


def _header_kind(cell: str) -> str | None:
    key = re.sub(r"[^a-zа-я0-9_]+", "", (cell or "").strip().lower())
    for needle, kind in CONTACT_HEADERS.items():
        if needle in key:
            return kind
    return None


def _rows_from_csv(text: str, dialect_hint: str | None = None) -> list[list[str]]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        if dialect_hint == "tsv" or "\t" in sample and sample.count("\t") > sample.count(","):
            dialect.delimiter = "\t"
    reader = csv.reader(StringIO(text), dialect)
    return [[(c or "").strip() for c in row] for row in reader if any((c or "").strip() for c in row)]


def _from_table(rows: list[list[str]]) -> ParseResult:
    result = ParseResult()
    if not rows:
        result.notes.append("таблица пустая")
        return result
    header = rows[0]
    kinds = [_header_kind(c) for c in header]
    has_header = any(kinds)
    data_rows = rows[1:] if has_header else rows
    if has_header:
        result.notes.append("нашёл заголовок, разбираю колонки")
        col_map = [(i, k) for i, k in enumerate(kinds) if k]
    else:
        col_map = []

    seen: set[tuple[str, str]] = set()
    for row in data_rows:
        picked: Classified | None = None
        extra_bits: list[str] = []
        if col_map:
            ranked: list[Classified] = []
            for idx, want in col_map:
                if idx >= len(row):
                    continue
                cell = row[idx]
                item = classify_token(cell)
                if not item:
                    continue
                if want != "any" and item.kind != want and item.kind != "unknown":
                    if want == "username" and item.kind == "user_id":
                        item.confidence *= 0.4
                    elif want == "phone" and item.kind != "phone":
                        item.confidence *= 0.3
                ranked.append(item)
            extra_bits = [c for i, c in enumerate(row) if c and all(i != idx for idx, _ in col_map)]
            if ranked:
                picked = max(ranked, key=lambda c: c.confidence)
        if picked is None:
            cells = [c for c in row if c]
            ranked = [classify_token(c) for c in cells]
            ranked = [c for c in ranked if c]
            if ranked:
                picked = max(ranked, key=lambda c: c.confidence)
                extra_bits = cells
        if picked is None or picked.kind == "unknown":
            joined = classify_line(" ".join(row))
            picked = joined
        if picked is None or picked.kind == "unknown":
            result.skipped += 1
            continue
        key = (picked.kind, picked.value.lower())
        if key in seen:
            result.skipped += 1
            continue
        seen.add(key)
        if extra_bits and not picked.extra:
            picked.extra = " ".join(extra_bits)[:240]
        picked.raw = picked.raw or " | ".join(row)[:400]
        result.contacts.append(picked)
    return result


def parse_json_text(text: str) -> ParseResult:
    result = ParseResult()
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        return parse_lines(text)
    items: list = []
    if isinstance(data, list):
        items = data
    elif isinstance(data, dict):
        for key in ("contacts", "users", "items", "data", "list"):
            if isinstance(data.get(key), list):
                items = data[key]
                break
        if not items:
            items = [data]
    seen: set[tuple[str, str]] = set()
    for item in items:
        if isinstance(item, str):
            got = classify_line(item)
        elif isinstance(item, (int, float)):
            got = classify_token(str(int(item)))
        elif isinstance(item, dict):
            blob = " ".join(
                str(item.get(k, ""))
                for k in (
                    "username",
                    "user",
                    "telegram",
                    "phone",
                    "mobile",
                    "id",
                    "user_id",
                    "tg_id",
                    "contact",
                )
                if item.get(k) not in (None, "")
            )
            got = classify_line(blob) if blob.strip() else None
            if got:
                got.extra = json.dumps(item, ensure_ascii=False)[:240]
        else:
            got = None
        if not got or got.kind == "unknown":
            result.skipped += 1
            continue
        key = (got.kind, got.value.lower())
        if key in seen:
            result.skipped += 1
            continue
        seen.add(key)
        result.contacts.append(got)
    result.notes.append("json")
    return result


def parse_lines(text: str) -> ParseResult:
    result = ParseResult()
    seen: set[tuple[str, str]] = set()
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        got = classify_line(line)
        if not got:
            if line.strip() and not line.strip().startswith("#"):
                result.skipped += 1
            continue
        key = (got.kind, got.value.lower())
        if key in seen:
            result.skipped += 1
            continue
        seen.add(key)
        result.contacts.append(got)
    return result


def parse_xlsx(data: bytes, sheet_name: str | None = None) -> ParseResult:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    if sheet_name and sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
    else:
        sheet = wb.active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        if any(cells):
            rows.append(cells)
    wb.close()
    result = _from_table(rows)
    result.notes.append(f"xlsx:{sheet.title}")
    return result


def parse_xlsx_all_sheets(data: bytes) -> dict[str, ParseResult]:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    out: dict[str, ParseResult] = {}
    for name in wb.sheetnames:
        sheet = wb[name]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append(cells)
        result = _from_table(rows)
        result.notes.append(f"xlsx:{name}")
        out[name] = result
    wb.close()
    return out


def parse_sql_text(text: str) -> ParseResult:
    result = ParseResult()
    seen: set[tuple[str, str]] = set()
    for line in text.replace("\r\n", "\n").split("\n"):
        line = line.strip()
        if not line or line.startswith("--") or line.startswith("/*"):
            continue
        for token in re.findall(r"'([^']*)'|\"([^\"]*)\"|(@?[A-Za-z0-9_+.-]+)", line):
            raw = next((t for t in token if t), "")
            if not raw:
                continue
            got = classify_token(raw) or classify_line(raw)
            if not got or got.kind == "unknown":
                continue
            key = (got.kind, got.value.lower())
            if key in seen:
                result.skipped += 1
                continue
            seen.add(key)
            result.contacts.append(got)
    result.notes.append("sql")
    return result


def parse_contacts_file(data: bytes, filename: str = "") -> ParseResult:
    name = (filename or "").lower()
    ext = Path(name).suffix.lower()
    if ext == ".sql":
        return parse_sql_text(decode_bytes(data))
    if ext in {".xlsx", ".xlsm"}:
        return parse_xlsx(data)
    text = decode_bytes(data)
    if ext == ".json" or (text.lstrip()[:1] in {"{", "["} and ext in {"", ".txt", ".json"}):
        stripped = text.lstrip()
        if stripped.startswith("{") or stripped.startswith("["):
            try:
                json.loads(stripped)
                return parse_json_text(stripped)
            except json.JSONDecodeError:
                pass
    if ext in {".csv", ".tsv"} or (ext == "" and ("," in text[:200] or ";" in text[:200])):
        delim = "tsv" if ext == ".tsv" else None
        rows = _rows_from_csv(text, delim)
        if rows and (len(rows[0]) > 1 or ext in {".csv", ".tsv"}):
            result = _from_table(rows)
            result.notes.append(ext or "csv")
            return result
    result = parse_lines(text)
    if ext:
        result.notes.append(ext.lstrip("."))
    return result
